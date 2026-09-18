"""Authenticated, temporary offer imports and as-of-list media reviews.

No imported workbook or report data is written to the static site or repository.
"""
import base64
import datetime as dt
import hashlib
import io
import json
import re
import zipfile

import offer_db as db
import offer_performance as performance

ASIN = re.compile(r"\bB[A-Z0-9]{9}\b", re.I)


def valid_date(value):
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(value or "")):
        raise ValueError("日期必须为 YYYY-MM-DD / Dates must use YYYY-MM-DD")
    return dt.date.fromisoformat(value).isoformat()


def catalog():
    historical = []
    for source in performance.catalog()["batches"]:
        historical.append({**source, "listDate": source.get("listDate") or source["launchDate"],
                           "logicalId": source["id"], "importedAt": None, "importedBy": None})
    return {"ok": True, "batches": historical}


def parse_workbook(body):
    from openpyxl import load_workbook
    name = str(body.get("sourceFile") or "").strip()
    if not name.lower().endswith(".xlsx") or len(name) > 200:
        raise ValueError("请选择 XLSX 文件 / Choose an XLSX file")
    date = valid_date(body.get("listDate"))
    try:
        raw = base64.b64decode(body.get("fileBase64") or "", validate=True)
    except (ValueError, TypeError) as exc:
        raise ValueError("Invalid workbook encoding") from exc
    if not raw or len(raw) > 2_500_000:
        raise ValueError("每个文件最多 2.5 MB / Maximum workbook size is 2.5 MB")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist()) > 1000 or sum(x.file_size for x in archive.infolist()) > 30_000_000:
                raise ValueError("Workbook expands beyond the supported size")
        book = load_workbook(io.BytesIO(raw), data_only=True)
    except (zipfile.BadZipFile, OSError, KeyError) as exc:
        raise ValueError("Invalid XLSX workbook") from exc
    offers, warnings, errors = {}, [], []
    def color(cell):
        fill = cell.fill.fgColor
        return f"{fill.type}:{fill.index}:{fill.tint}" if cell.fill.patternType else ""
    for sheet in sorted(book, key=lambda s: s.title.lower() != "list of offers"):
        if sheet.max_row > 20000 or sheet.max_column > 100:
            raise ValueError("Workbook sheet is too large")
        rows = list(sheet.iter_rows())
        header = next((i for i, row in enumerate(rows[:100]) if any(re.fullmatch(r"merchant\s*id|商家\s*id", str(c.value or "").strip(), re.I) for c in row)), None)
        if header is None:
            continue
        legend = {}
        for row in rows[:header]:
            label = " ".join(str(c.value or "") for c in row).strip()
            if label:
                for c in row:
                    if color(c):
                        legend[color(c)] = label
        headers = [str(c.value or "").strip().lower() for c in rows[header]]
        def field(row, pattern):
            index = next((i for i, h in enumerate(headers) if re.search(pattern, h)), None)
            return str(row[index].value or "").strip() if index is not None else ""
        for row in rows[header + 1:]:
            if not any(c.value is not None for c in row):
                continue
            mid = field(row, r"^(merchant\s*id|商家\s*id)$")
            if not re.fullmatch(r"[1-9]\d{0,12}", mid):
                errors.append(f"{sheet.title}:{row[0].row} 无效 Merchant ID / Invalid Merchant ID")
                continue
            merchant = field(row, r"^(merchant\s*name|商家名(?:称)?)$")
            if not merchant:
                errors.append(f"{sheet.title}:{row[0].row} 缺少商家名称 / Missing merchant name")
            notes = field(row, r"^(推荐信息|推荐理由|recommendation|recommendation reason|reason|notes)$")
            products = sorted(set(ASIN.findall(field(row, r"brand product list|asin|单品列表").upper())))
            # Negated sentences are never silently turned into positive targets.
            positive, ambiguous = [], []
            for sentence in re.split(r"[\n。;；]", notes):
                found = ASIN.findall(sentence.upper())
                (ambiguous if re.search(r"不推|不要|不建议|无需|避免|暂停|停止|not\b|avoid\b|stop\b", sentence, re.I) else positive).extend(found)
            priority = field(row, r"^(priority|等级|推荐等级)$") or next((legend[color(c)] for c in row[:3] if color(c) in legend), "未识别 / Unknown")
            aov_raw = field(row, r"aov|客单价")
            match = re.search(r"\d+(?:\.\d+)?", aov_raw.replace(",", ""))
            aov = float(match.group()) if match and not re.search(r"EUR|GBP|CNY|€|£|¥", aov_raw, re.I) else None
            offer = {"merchantId": mid, "merchantName": merchant[:160], "category": field(row, r"category|品类"),
                     "asins": sorted(set(products + positive)), "productAsins": products, "reasonAsins": sorted(set(positive)),
                     "ambiguousAsins": sorted(set(ambiguous)), "notes": notes, "priority": priority,
                     "referenceAov": aov, "aovRaw": aov_raw,
                     "sourceRows": [{"sheet": sheet.title, "row": row[0].row, "fill": color(row[0]), "aovRaw": aov_raw,
                                     "values": {headers[i]: str(c.value or "") for i, c in enumerate(row) if headers[i]}}]}
            if mid in offers:
                previous = offers[mid]
                same_sheet = any(x["sheet"] == sheet.title for x in previous["sourceRows"])
                if (same_sheet and any(previous[k] != offer[k] for k in ("asins", "referenceAov", "priority"))) or any(previous[k] and offer[k] and previous[k] != offer[k] for k in ("merchantName", "notes")):
                    errors.append(f"Merchant ID {mid} 同文件内容冲突 / Conflicting duplicate rows")
                if not same_sheet and previous["referenceAov"] and offer["referenceAov"] and previous["referenceAov"] != offer["referenceAov"]:
                    warnings.append(f"{mid}: 两表 AOV 不同；分层采用 Offer 主表值 {previous['referenceAov']}，产品表值 {offer['referenceAov']} 保留在来源记录 / Different sheet AOVs; offer-sheet value is used")
                for k in ("asins", "productAsins", "reasonAsins", "ambiguousAsins"):
                    previous[k] = sorted(set(previous[k] + offer[k]))
                for k in ("notes", "referenceAov", "aovRaw", "category"):
                    previous[k] = previous[k] or offer[k]
                if previous["priority"] == "未识别 / Unknown":
                    previous["priority"] = offer["priority"]
                previous["sourceRows"] += offer["sourceRows"]
            else:
                offers[mid] = offer
            if ambiguous:
                warnings.append(f"{mid}: 理由中有否定句 ASIN，未自动纳入理由目标 / Negated ASIN excluded from reason targets")
            if priority == "未识别 / Unknown":
                warnings.append(f"{mid}: 未识别原表等级 / Unrecognized source priority")
    book.close()
    if not 1 <= len(offers) <= 200:
        errors.append("每份清单需包含 1–200 个商家 / Each list requires 1–200 merchants")
    return {"name": name.rsplit(".", 1)[0], "sourceFile": name, "listDate": date, "launchDate": date,
            "fileHash": hashlib.sha256(raw).hexdigest(), "parserVersion": 1,
            "offers": list(offers.values()), "warnings": warnings, "errors": errors}


def import_workbook(body, actor):
    parsed = parse_workbook(body)
    if body.get("action") == "review-preview":
        return {"ok": True, "batch": parsed}
    if parsed["errors"]:
        raise ValueError("; ".join(parsed["errors"][:8]))
    logical_id = str(body.get("logicalId") or "").strip()
    if len(logical_id) > 160:
        raise ValueError("Invalid logical list ID")
    identity = "session-" + parsed["fileHash"]
    return {"ok": True, "batch": {**parsed, "id": identity, "logicalId": logical_id or identity,
            "importedAt": dt.datetime.now(dt.timezone.utc).isoformat(), "importedBy": actor[:160], "temporary": True}}


def temporary_batches(raw):
    """Validate browser-session input before using it to build SQL report IDs."""
    if not isinstance(raw, list) or len(raw) > 30:
        raise ValueError("本次会话最多 30 份临时清单 / At most 30 temporary lists")
    output, seen = [], set()
    for batch in raw:
        if not isinstance(batch, dict) or not re.fullmatch(r"session-[a-f0-9]{64}", str(batch.get("id") or "")):
            raise ValueError("Invalid temporary list identity")
        if batch["id"] in seen:
            continue
        seen.add(batch["id"])
        valid_date(batch.get("listDate"))
        if not isinstance(batch.get("name"), str) or not isinstance(batch.get("offers"), list) or not 1 <= len(batch["offers"]) <= 200:
            raise ValueError("Invalid temporary list")
        ids = set()
        for offer in batch["offers"]:
            if not isinstance(offer, dict) or not re.fullmatch(r"[1-9]\d{0,12}", str(offer.get("merchantId") or "")) or offer["merchantId"] in ids:
                raise ValueError("Invalid or duplicate Merchant ID")
            ids.add(offer["merchantId"])
            if not isinstance(offer.get("merchantName"), str) or not offer["merchantName"].strip() or not isinstance(offer.get("asins"), list):
                raise ValueError("Invalid merchant fields")
            if any(not isinstance(a, str) or not ASIN.fullmatch(a) for a in offer["asins"]):
                raise ValueError("Invalid recommended ASIN")
        output.append(batch)
    return output


def resolve_cohort(batches, selected_ids, recommendation):
    valid_date(recommendation)
    if not isinstance(selected_ids, list) or not 1 <= len(selected_ids) <= 30 or any(not isinstance(x, str) for x in selected_ids):
        raise ValueError("请选择 1–30 份清单 / Choose 1–30 lists")
    selected = [b for b in batches if b["id"] in selected_ids]
    if set(selected_ids) != {b["id"] for b in selected}:
        raise ValueError("清单已失效，请重新加载 / Unknown list; reload catalog")
    excluded = [b["id"] for b in selected if b["listDate"] > recommendation]
    by_merchant = {}
    for batch in sorted(selected, key=lambda b: b["listDate"]):
        if batch["id"] in excluded:
            continue
        for offer in batch["offers"]:
            mid = offer["merchantId"]
            current = by_merchant.get(mid)
            if current and current["listDate"] == batch["listDate"]:
                signature = lambda o: json.dumps({k: o.get(k) for k in ("asins", "notes", "priority", "referenceAov")}, sort_keys=True)
                if signature(current) != signature(offer):
                    raise ValueError(f"{mid} 同日版本冲突，请只选一个版本 / Conflicting same-day versions; select one")
            by_merchant[mid] = {**offer, "listDate": batch["listDate"], "batchId": batch["id"], "listName": batch["name"]}
    if not 1 <= len(by_merchant) <= 200:
        raise ValueError("有效清单需包含 1–200 个去重商家 / Select 1–200 eligible merchants")
    for mid, offer in by_merchant.items():
        history = [b for b in batches if any(o["merchantId"] == mid for o in b["offers"])]
        offer["importCount"] = sum(bool(b.get("importedAt")) for b in history)
        offer["listCount"] = len({b.get("logicalId", b["id"]) for b in history})
        offer["history"] = [{k: b.get(k) for k in ("id", "name", "listDate", "importedAt", "importedBy")} for b in history]
    return list(by_merchant.values()), excluded


def review(body):
    batches = catalog()["batches"] + temporary_batches(body.get("temporaryBatches", []))
    recommendation = body.get("launchDate")
    offers, excluded = resolve_cohort(batches, body.get("batchIds"), recommendation)
    window = performance.date_window(recommendation, body.get("startDate"), body.get("endDate"), body.get("beforeStart"), body.get("beforeEnd"))
    query = {k: [v] for k, v in {"launchDate": recommendation, "startDate": window["startDate"], "endDate": window["endDate"],
             "beforeStart": window["beforeStart"], "beforeEnd": window["beforeEnd"],
             "merchantIds": ",".join(o["merchantId"] for o in offers), "currency": "USD", "skipHistory": "1"}.items()}
    selected = str(body.get("merchantId") or "")
    if selected:
        query["merchantId"] = [selected]
        report = performance.report(query)
    else:
        part = body.get("part", "summary")
        if part not in ("summary", "relations"):
            raise ValueError("Invalid review part")
        # Independent queries avoid serial summary/detail latency.
        report = performance.report({**query, **({"action": ["relations"]} if part == "relations" else {})})
    report.update({"offers": offers, "excludedBatchIds": excluded, "currency": "USD", "recommendationDate": recommendation,
                   "dataNote": "Latest record date is not proof of ingestion completeness; clicks are not published post counts."})
    return report


def handle_request(target, method, query):
    from auth import _read_json_body, current_user_for_target
    try:
        if method == "POST" and db.first_query_value(query, "action"):
            db.send_json(target, 405, {"ok": False, "error": "Method not allowed"})
            return
        if method == "GET":
            db.send_json(target, 200, catalog())
            return
        body = _read_json_body(target, max_size=3_600_000)
        if not isinstance(body, dict):
            raise ValueError("Expected a JSON object")
        action = body.get("action")
        if action == "review":
            result = review(body)
        elif action in {"review-preview", "review-import"}:
            user = current_user_for_target(target) or {}
            result = import_workbook(body, str(user.get("username") or user.get("id") or "local-development"))
        else:
            raise ValueError("Unsupported review action")
        db.send_json(target, 200, result)
    except (ValueError, TypeError, json.JSONDecodeError) as error:
        db.send_json(target, 400, {"ok": False, "error": str(error)})
    except Exception as error:
        db.send_db_error(target, error)
