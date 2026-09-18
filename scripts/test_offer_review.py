"""Synthetic import, temporal cohort and authenticated route regression tests."""
import base64
import io
import json
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import offer_review as review
import offer_performance as performance
from test_vercel_db_wsgi import load_app_module, request
from openpyxl import Workbook
from openpyxl.styles import PatternFill


def workbook():
    book = Workbook()
    sheet = book.active
    sheet.title = "List of Offers"
    sheet.append([None, "高优先级 offer"])
    sheet.cell(1, 1).fill = PatternFill('solid', fgColor='D6EEDD')
    sheet.append(['Merchant ID', 'Merchant Name', '推荐信息', 'AOV'])
    sheet.append([101, 'Example', '推荐 B012345678；不要推 B000000000', '$129'])
    sheet.cell(3, 1).fill = PatternFill('solid', fgColor='D6EEDD')
    products = book.create_sheet('Brand Product List')
    products.append(['Merchant ID', 'Merchant Name', 'Top ASINs'])
    products.append([101, 'Example', 'B098765432'])
    buffer = io.BytesIO()
    book.save(buffer)
    return {'sourceFile': 'Fixture (9-7).xlsx', 'fileBase64': base64.b64encode(buffer.getvalue()).decode(), 'listDate': '2026-09-07'}


class ReviewTests(unittest.TestCase):
    def test_import_joins_complementary_sheets_and_preserves_evidence(self):
        parsed = review.parse_workbook(workbook())
        self.assertEqual(parsed['errors'], [])
        self.assertEqual(len(parsed['offers']), 1)
        offer = parsed['offers'][0]
        self.assertEqual(offer['asins'], ['B012345678', 'B098765432'])
        self.assertEqual(offer['productAsins'], ['B098765432'])
        self.assertEqual(offer['reasonAsins'], ['B012345678'])
        self.assertEqual(offer['ambiguousAsins'], ['B000000000'])
        self.assertEqual(offer['referenceAov'], 129)
        self.assertIn('高优先级', offer['priority'])
        self.assertEqual(len(offer['sourceRows']), 2)
        self.assertNotIn('importedAt', parsed)

    def batches(self):
        offer = {'merchantId': '101', 'merchantName': 'Example', 'asins': ['B012345678'], 'notes': '', 'referenceAov': 10}
        return [dict(id='old', logicalId='old', listDate='2026-09-01', name='old', offers=[offer], importedAt=None),
                dict(id='new', logicalId='new', listDate='2026-09-07', name='new', offers=[{**offer, 'asins': ['B098765432']}], importedAt='2026-09-18T00:00:00Z'),
                dict(id='future', logicalId='future', listDate='2026-09-14', name='future', offers=[{**offer, 'merchantId': '1101'}], importedAt=None)]

    def test_future_versions_excluded_exact_ids_and_latest_targets(self):
        offers, excluded = review.resolve_cohort(self.batches(), ['old', 'new', 'future'], '2026-09-07')
        self.assertEqual(excluded, ['future'])
        self.assertEqual([x['merchantId'] for x in offers], ['101'])
        self.assertEqual(offers[0]['asins'], ['B098765432'])
        self.assertEqual(offers[0]['importCount'], 1)
        self.assertEqual(offers[0]['listCount'], 2)

    def test_same_day_conflict_blocks_and_unknown_lists_fail(self):
        batches = self.batches()
        batches[1]['listDate'] = batches[0]['listDate']
        with self.assertRaisesRegex(ValueError, 'Conflicting'):
            review.resolve_cohort(batches, ['old', 'new'], '2026-09-07')
        with self.assertRaisesRegex(ValueError, 'Unknown list'):
            review.resolve_cohort(batches, ['missing'], '2026-09-07')

    def test_retry_is_idempotent_and_server_controls_import_actor_time(self):
        body = {**workbook(), 'action': 'review-import', 'importedAt': '1990-01-01'}
        with patch.object(review.db, 'db_connection') as connection:
            result = review.import_workbook(body, 'authenticated-user')['batch']
            retry = review.import_workbook(body, 'authenticated-user')['batch']
            connection.assert_not_called()
        self.assertEqual(result['id'], retry['id'])
        self.assertEqual(result['importedBy'], 'authenticated-user')
        self.assertNotEqual(result['importedAt'], '1990-01-01')
        self.assertTrue(result['temporary'])
        self.assertEqual(len(review.temporary_batches([result, retry])), 1)

    def test_temporary_report_input_rejects_invalid_identity_before_query(self):
        batch = review.import_workbook({**workbook(), 'action': 'review-import'}, 'fixture')['batch']
        batch['offers'][0]['merchantId'] = '101 OR 1=1'
        with patch.object(review.db, 'db_connection') as connection, self.assertRaises(ValueError):
            review.review({'temporaryBatches': [batch]})
        connection.assert_not_called()

    def test_preview_never_opens_database(self):
        with patch.object(review.db, 'db_connection') as conn:
            self.assertTrue(review.import_workbook({**workbook(), 'action': 'review-preview'}, 'reader')['ok'])
            conn.assert_not_called()

    def test_usd_filter_and_record_click_asin(self):
        with patch.object(performance.db, 'fetch_all', return_value=[]) as fetch:
            performance._read(None, 'cnpscy_amazon_order', {'advert_id', 'order_time_day', 'amount', 'currency'}, ['101'], '2026-09-01', '2026-09-14', performance.ORDER_FIELDS, known_watermark='2026-09-14', currency='USD')
            self.assertIn("`currency` = 'USD'", fetch.call_args.args[1])
            performance._read(None, 'cnpscy_amazon_click', {'advert_id', 'time_day', 'click', 'asin'}, ['101'], '2026-09-01', '2026-09-14', {'clicks': ['click']}, detail=True, known_watermark='2026-09-14')
            self.assertIn('r.`asin`', fetch.call_args.args[1])
        self.assertEqual(performance.target_identity({'target_asin': 'B012345678', 'link_type': 'storefront'}), ('storefront', ''))

    def test_media_zeros_do_not_assign_clicks_to_purchased_asins(self):
        supported = dict.fromkeys(performance.METRICS, True)
        media, links = performance.summarize_details([{'merchantId': '101', 'publisherId': '7', 'day': '20260907', 'revenue': 10, 'purchasedAsin': 'B012345678'}], performance.date_window('2026-09-07'), supported, '2026-09-14')
        self.assertEqual(media[0]['before']['clicks'], 0)
        self.assertIsNone(links[0]['after']['clicks'])

    def test_wsgi_preview_authentication_and_no_public_data(self):
        app = load_app_module()
        with patch.dict('os.environ', {'OI_AUTH_ENABLED': '0', 'VERCEL_ENV': '', 'VERCEL': ''}):
            response = request(app.app, 'ui-offer-performance', method='POST', body={**workbook(), 'action': 'review-preview'})
            self.assertEqual(response['status'], 200)
            self.assertEqual(response['headers']['Cache-Control'], 'no-store')
        with patch.dict('os.environ', {'OI_AUTH_ENABLED': '1', 'OI_SESSION_SECRET': 'test-secret'}), patch.object(review, 'parse_workbook') as parse:
            response = request(app.app, 'ui-offer-performance', method='POST', body={**workbook(), 'action': 'review-import'})
            self.assertIn(response['status'], [401, 503])
            parse.assert_not_called()


if __name__ == '__main__':
    unittest.main()
